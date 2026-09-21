import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_blueprint_docx():
    doc = docx.Document()

    # Set standard academic margins
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1.25)
        s.right_margin = Inches(1)

    with open('Dokumen_Cetak_Biru_Arsitektur_WebGIS_DIPA2026.md', 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code = False
    code_buffer = []

    for line in lines:
        raw = line.rstrip('\n')
        
        if raw.startswith('```'):
            if in_code:
                in_code = False
                p = doc.add_paragraph('\n'.join(code_buffer))
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                for run in p.runs:
                    run.font.name = 'Courier New'
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(40, 40, 40)
                code_buffer = []
            else:
                in_code = True
                code_buffer = []
            continue

        if in_code:
            code_buffer.append(raw)
            continue

        if raw.startswith('# '):
            p = doc.add_paragraph()
            run = p.add_run(raw[2:])
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(15, 32, 67)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        elif raw.startswith('## '):
            p = doc.add_paragraph()
            run = p.add_run(raw[3:])
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(30, 60, 114)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
        elif raw.startswith('### '):
            p = doc.add_paragraph()
            run = p.add_run(raw[4:])
            run.bold = True
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(45, 52, 54)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
        elif raw.startswith('#### '):
            p = doc.add_paragraph()
            run = p.add_run(raw[5:])
            run.bold = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(50, 50, 50)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
        elif raw.startswith('---'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run('—' * 50)
            run.font.color.rgb = RGBColor(180, 180, 180)
        elif raw.startswith('|') and '---' in raw:
            continue
        elif raw.startswith('|'):
            p = doc.add_paragraph(raw)
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
        elif raw.startswith('- ') or raw.startswith('* '):
            p = doc.add_paragraph(raw[2:], style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                run.font.size = Pt(10.5)
        elif raw.strip():
            p = doc.add_paragraph(raw)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            for run in p.runs:
                run.font.size = Pt(10.5)

    out_path = 'Dokumen_Cetak_Biru_Arsitektur_WebGIS_DIPA2026.docx'
    doc.save(out_path)
    print(f'Successfully generated {out_path}!')

if __name__ == '__main__':
    create_blueprint_docx()
