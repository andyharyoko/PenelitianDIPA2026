import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_bab2():
    doc = Document()

    # Kertas dan Margin A4 (Standar Jurnal/Skripsi)
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(4)
        section.right_margin = Cm(3)
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(3)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5

    # Judul
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run('BAB II\nTINJAUAN PUSTAKA')
    r_title.bold = True
    
    doc.add_paragraph()

    # 2.1 Teori Pendukung
    doc.add_paragraph('2.1 Teori Pendukung', style='Normal').runs[0].bold = True
    p1 = doc.add_paragraph('Penelitian ini dilandasi oleh beberapa kerangka teoretis utama, yaitu Sistem Pendukung Keputusan (SPK), Sistem Informasi Geografis (SIG), dan optimasi kombinatorial matematis. Menurut Turban et al. (2005), SPK adalah sistem berbasis komputer interaktif yang membantu pengambil keputusan memanfaatkan data dan model untuk menyelesaikan masalah tidak terstruktur dan semi-terstruktur. Dalam perkembangannya, model Multi-Criteria Decision Making (MCDM) seperti Simple Additive Weighting (SAW) dan AHP-TOPSIS menjadi pilar inti dalam mengkuantifikasi parameter benefit dan cost dari berbagai ragam alternatif keputusan (Hwang & Yoon, 1981).')
    p1.paragraph_format.first_line_indent = Cm(1.27)

    p2 = doc.add_paragraph('Untuk merepresentasikan keputusan di bidang spasial, konsep Sistem Informasi Geografis (SIG) digunakan sebagai medium pemrosesan atribut keruangan. Longley et al. (2015) mendefinisikan SIG sebagai teknologi multidisiplin yang tidak hanya memetakan fenomena bumi, tetapi juga mengintegrasikan berbagai layer informasi geografis (seperti jarak, batas infrastruktur, dan titik fasilitas umum) untuk menunjang analisis keruangan. Pengambilan data penelitian melalui OpenStreetMap (OSM) secara teoretis bertumpu pada landasan Volunteered Geographic Information (VGI) yang secara eksponensial terus berkembang (Haklay & Weber, 2008).')
    p2.paragraph_format.first_line_indent = Cm(1.27)

    p3 = doc.add_paragraph('Di sisi lain, proses perencanaan evaluasi lapangan (ground truth) pada penelitian ini mengacu pada teori Riset Operasi, secara khusus model Traveling Salesperson Problem (TSP). TSP bertujuan untuk secara matematis menemukan rute sirkuit terpendek yang mengunjungi sekumpulan titik koordinat spasial tepat satu kali, yang merupakan representasi fundamental dari algoritma komputasi NP-Hard dalam optimasi logistik (Applegate et al., 2006).')
    p3.paragraph_format.first_line_indent = Cm(1.27)

    doc.add_paragraph()

    # 2.2 State of the Art
    doc.add_paragraph('2.2 Kajian Literatur (State of the Art) dan Novelty', style='Normal').runs[0].bold = True
    p4 = doc.add_paragraph('Penelitian yang berfokus pada optimasi penentuan lokasi (site selection) telah dieksplorasi secara masif dalam kurun waktu lima tahun terakhir. Studi terkini oleh Yaman (2024) dan Yousefi et al. (2024) membuktikan bahwa penggabungan GIS dan MCDM (GIS-MCDM) merupakan pendekatan paling efektif dalam memetakan kelayakan lahan berdasarkan parameter ekonomi, infrastruktur, dan batasan demografis. Selain itu, Ccatamayo-Barrios et al. (2023) secara empiris mengkomparasikan efisiensi metode AHP dan TOPSIS, dengan konklusi bahwa metode TOPSIS menawarkan tingkat konsistensi peringkat yang lebih reliabel dalam menghadapi alternatif keputusan yang memiliki simpangan jarak ekstrem terhadap kriteria solusi ideal.')
    p4.paragraph_format.first_line_indent = Cm(1.27)

    p5 = doc.add_paragraph('Namun demikian, berdasarkan tinjauan literatur (State of the Art), mayoritas studi GIS-MCDM kontemporer (Yaman, 2024; Yousefi et al., 2024; Ccatamayo-Barrios et al., 2023) memiliki kesenjangan metodologis (research gap). Kebanyakan dari riset tersebut mengaplikasikan kalkulasi pembobotan secara langsung terhadap ratusan hingga ribuan titik lokasi koordinat mentah tanpa tahapan penyaringan (filtering) awal. Pendekatan ini rentan memicu inefisiensi komputasi serta memperbesar peluang terjadinya fenomena "rank reversal" akibat outlier. Oleh karena itu, novelty (kebaruan) pada penelitian ini adalah intervensi lapisan pra-pemrosesan menggunakan algoritma Pareto Optimality (Anysz et al., 2021) untuk secara absolut membuang alternatif yang terdominasi sebelum masuk ke tahap eksekusi MCDM. Kebaruan kedua diwujudkan melalui pengintegrasian optimasi rute TSP (Kim et al., 2025) ke dalam lokasi-lokasi hasil saringan Pareto guna menciptakan tata panduan validasi lapangan (ground truth) yang secara kuantitatif menekan biaya operasional survei.')
    p5.paragraph_format.first_line_indent = Cm(1.27)

    doc.add_paragraph()

    # 2.3 Roadmap
    doc.add_paragraph('2.3 Peta Jalan (Roadmap) Penelitian', style='Normal').runs[0].bold = True
    p6 = doc.add_paragraph('Guna menjamin kontinuitas inovasi serta dampak akademis jangka panjang dari model sistem penunjang keputusan spasial yang dibangun, peta jalan (roadmap) penelitian ini direkayasa ke dalam tiga lintasan capaian utama yang berkelanjutan:')
    
    plans = [
        "Fase 1 (Tahap Saat Ini): Perancangan kerangka algoritmik sistem pendukung keputusan melalui akuisisi data geospasial OpenStreetMap yang digabungkan dengan penyaringan Pareto Optimality, evaluasi komparatif SAW dan AHP-TOPSIS, serta optimasi efisiensi logistik lapangan via simulasi rute TSP.",
        "Fase 2 (Jangka Menengah - 1 hingga 2 Tahun Kedepan): Transformasi model arsitektur menjadi purwarupa (prototype) WebGIS interaktif berskala produksi berbasis Cloud Computing. Platform ini akan mendukung fitur dynamic-MCDM, di mana para pemangku kepentingan dapat melakukan rekalibrasi bobot parameter kriteria secara real-time dan melihat perubahan rekomendasi lokasi kafe berdasar sentimen ekonomi saat itu.",
        "Fase 3 (Jangka Panjang - 3 Tahun Kedepan): Pengintegrasian model kecerdasan buatan (Machine Learning) seperti Random Forest atau Deep Neural Networks untuk menganalisis dan mendeteksi pola historis keberhasilan bisnis komersial di kawasan tertentu. Fase ini akan mengubah paradigma sistem dari yang sebelumnya hanya bersifat preskriptif menjadi prediktif, di mana sistem mampu mengestimasi persentase profitabilitas/ROI suatu lokasi sebelum dibangun."
    ]

    for plan in plans:
        p_plan = doc.add_paragraph(plan, style='List Number')
        p_plan.paragraph_format.line_spacing = 1.5

    doc.save('Bab_II_Tinjauan_Pustaka.docx')
    print("Dokumen Bab II berhasil dibuat.")

if __name__ == "__main__":
    create_bab2()
